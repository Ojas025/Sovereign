"""Document generation tools: generate_docx and generate_pdf.

Allows the agent to output structured deliverables (Word and PDF documents)
from markdown text. PathJail is enforced for workspace confinement.
"""

from __future__ import annotations

from collections.abc import Mapping
from pathlib import Path

from workbench.config import ToolsConfig
from workbench.core.protocols import ToolContext, ToolResult
from workbench.tools.jail import PathJail, PathJailError

_PATH_JAIL = "path_jail"


class _BadArgument(Exception):
    """A model-supplied argument failed structural validation."""


def _require_str(arguments: Mapping[str, object], name: str, *, allow_empty: bool = False) -> str:
    value = arguments.get(name)
    if not isinstance(value, str) or (not allow_empty and not value):
        expectation = "a string" if allow_empty else "a non-empty string"
        raise _BadArgument(f"{name} must be {expectation}")
    return value


def _jail(tools: ToolsConfig, context: ToolContext) -> PathJail:
    """Per-call jail: workspace from call context, extra read roots from config."""
    return PathJail(context.workspace_root, [Path(root) for root in tools.read_roots])


def _build_docx(title: str, content: str, output_path: Path) -> int:
    import docx
    from docx.shared import Pt, RGBColor

    from workbench.tools.markdown_parser import (
        CodeBlock,
        Heading,
        HorizontalRule,
        ListBlock,
        Paragraph,
        Table,
        parse_inlines,
        parse_markdown,
    )

    doc = docx.Document()
    if title:
        doc.add_heading(title, level=0)

    blocks = parse_markdown(content)
    for block in blocks:
        if isinstance(block, Heading):
            lvl = min(max(block.level, 1), 9)
            doc.add_heading(block.text, level=lvl)
        elif isinstance(block, Paragraph):
            p = doc.add_paragraph()
            spans = parse_inlines(block.text)
            for span in spans:
                run = p.add_run(span.text)
                if span.bold:
                    run.bold = True
                if span.italic:
                    run.italic = True
                if span.code:
                    run.font.name = "Courier New"
        elif isinstance(block, ListBlock):
            for item in block.items:
                style = "List Number" if block.ordered else "List Bullet"
                p = doc.add_paragraph(style=style)
                spans = parse_inlines(item)
                for span in spans:
                    run = p.add_run(span.text)
                    if span.bold:
                        run.bold = True
                    if span.italic:
                        run.italic = True
                    if span.code:
                        run.font.name = "Courier New"
        elif isinstance(block, CodeBlock):
            p = doc.add_paragraph()
            run = p.add_run(block.code)
            run.font.name = "Courier New"
            run.font.size = Pt(9.5)
        elif isinstance(block, Table):
            if block.headers or block.rows:
                cols = max(len(block.headers), max((len(r) for r in block.rows), default=0))
                rows_count = (1 if block.headers else 0) + len(block.rows)
                tbl = doc.add_table(rows=rows_count, cols=cols)
                tbl.style = "Table Grid"
                cur_row = 0
                if block.headers:
                    for c_idx, h in enumerate(block.headers):
                        cell = tbl.rows[cur_row].cells[c_idx]
                        cell.text = h
                        for p in cell.paragraphs:
                            for r in p.runs:
                                r.bold = True
                    cur_row += 1
                for row_data in block.rows:
                    for c_idx, val in enumerate(row_data):
                        if c_idx < cols:
                            tbl.rows[cur_row].cells[c_idx].text = val
                    cur_row += 1
        elif isinstance(block, HorizontalRule):
            p = doc.add_paragraph()
            r = p.add_run("―" * 32)
            r.font.color.rgb = RGBColor(180, 180, 180)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    return output_path.stat().st_size


def _build_pdf(title: str, content: str, output_path: Path) -> int:
    from fpdf import FPDF

    from workbench.tools.markdown_parser import (
        CodeBlock,
        Heading,
        HorizontalRule,
        ListBlock,
        Paragraph,
        Table,
        parse_markdown,
    )

    def _sanitize(text: str) -> str:
        replacements = {
            "—": "--",
            "–": "-",
            "“": '"',
            "”": '"',
            "‘": "'",
            "’": "'",
            "•": "*",
            "…": "...",
            "\u200b": "",
            "\u00a0": " ",
            "™": "(TM)",
            "©": "(C)",
            "®": "(R)",
            "\t": "    ",
        }
        for k, v in replacements.items():
            text = text.replace(k, v)
        return text.encode("latin-1", "replace").decode("latin-1")

    pdf = FPDF(format="A4")
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()

    if title:
        pdf.set_font("Helvetica", "B", 18)
        pdf.cell(0, 12, _sanitize(title), new_x="LMARGIN", new_y="NEXT")
        pdf.ln(3)

    blocks = parse_markdown(content)
    for block in blocks:
        if isinstance(block, Heading):
            sizes = {1: 15, 2: 13, 3: 11, 4: 10, 5: 9, 6: 8}
            size = sizes.get(block.level, 11)
            pdf.ln(2)
            pdf.set_font("Helvetica", "B", size)
            pdf.cell(0, 8, _sanitize(block.text), new_x="LMARGIN", new_y="NEXT")
            pdf.ln(1)
        elif isinstance(block, Paragraph):
            pdf.set_font("Helvetica", "", 10)
            clean_text = _sanitize(block.text)
            pdf.multi_cell(0, 5, clean_text, new_x="LMARGIN", new_y="NEXT")
            pdf.ln(2)
        elif isinstance(block, ListBlock):
            pdf.set_font("Helvetica", "", 10)
            for i, item in enumerate(block.items, start=1):
                bullet = f"{i}. " if block.ordered else "- "
                clean_item = _sanitize(f"{bullet}{item}")
                pdf.multi_cell(0, 5, clean_item, new_x="LMARGIN", new_y="NEXT")
            pdf.ln(2)
        elif isinstance(block, CodeBlock):
            pdf.set_font("Courier", "", 9)
            pdf.set_fill_color(245, 245, 245)
            clean_code = _sanitize(block.code)
            for code_line in clean_code.splitlines():
                pdf.cell(0, 4.5, f"  {code_line}", fill=True, new_x="LMARGIN", new_y="NEXT")
            pdf.ln(2)
        elif isinstance(block, Table):
            if block.headers or block.rows:
                pdf.set_font("Helvetica", "", 9)
                with pdf.table() as tbl:
                    if block.headers:
                        h_row = tbl.row()
                        for h in block.headers:
                            h_row.cell(_sanitize(h))
                    for r_data in block.rows:
                        row = tbl.row()
                        for cell_val in r_data:
                            row.cell(_sanitize(cell_val))
                pdf.ln(2)
        elif isinstance(block, HorizontalRule):
            y = pdf.get_y() + 2
            pdf.set_draw_color(200, 200, 200)
            pdf.line(pdf.l_margin, y, pdf.w - pdf.r_margin, y)
            pdf.set_draw_color(0, 0, 0)
            pdf.set_y(y + 3)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    pdf.output(str(output_path))
    return output_path.stat().st_size


class GenerateDocxTool:
    name: str = "generate_docx"
    description: str = (
        "Generate a Microsoft Word (.docx) document from markdown content. "
        "The content is parsed as markdown and converted to a styled Word document."
    )
    parameters: Mapping[str, object] = {
        "type": "object",
        "properties": {
            "title": {"type": "string", "description": "Document title"},
            "content": {"type": "string", "description": "Document content in markdown format"},
            "path": {
                "type": "string",
                "description": "Output file path (workspace-relative, must end in .docx)",
            },
        },
        "required": ["title", "content", "path"],
    }

    def __init__(self, tools: ToolsConfig) -> None:
        self._tools = tools

    async def execute(
        self, arguments: Mapping[str, object], context: ToolContext
    ) -> ToolResult:
        try:
            raw_path = _require_str(arguments, "path")
            if not raw_path.lower().endswith(".docx"):
                raise _BadArgument("path must end with .docx")
            title = _require_str(arguments, "title", allow_empty=True)
            content = _require_str(arguments, "content", allow_empty=True)
            path = _jail(self._tools, context).resolve_write(raw_path)
        except _BadArgument as exc:
            return ToolResult(str(exc), is_error=True)
        except PathJailError as exc:
            return ToolResult(f"blocked: {exc}", is_error=True, blocked_reason=_PATH_JAIL)

        try:
            size = _build_docx(title, content, path)
        except Exception as exc:
            return ToolResult(f"failed to generate docx at {raw_path}: {exc}", is_error=True)

        return ToolResult(f"generated docx ({size} bytes) at {raw_path}")


class GeneratePdfTool:
    name: str = "generate_pdf"
    description: str = (
        "Generate a PDF document from markdown content. "
        "The content is parsed as markdown and rendered to a styled PDF."
    )
    parameters: Mapping[str, object] = {
        "type": "object",
        "properties": {
            "title": {"type": "string", "description": "Document title"},
            "content": {"type": "string", "description": "Document content in markdown format"},
            "path": {
                "type": "string",
                "description": "Output file path (workspace-relative, must end in .pdf)",
            },
        },
        "required": ["title", "content", "path"],
    }

    def __init__(self, tools: ToolsConfig) -> None:
        self._tools = tools

    async def execute(
        self, arguments: Mapping[str, object], context: ToolContext
    ) -> ToolResult:
        try:
            raw_path = _require_str(arguments, "path")
            if not raw_path.lower().endswith(".pdf"):
                raise _BadArgument("path must end with .pdf")
            title = _require_str(arguments, "title", allow_empty=True)
            content = _require_str(arguments, "content", allow_empty=True)
            path = _jail(self._tools, context).resolve_write(raw_path)
        except _BadArgument as exc:
            return ToolResult(str(exc), is_error=True)
        except PathJailError as exc:
            return ToolResult(f"blocked: {exc}", is_error=True, blocked_reason=_PATH_JAIL)

        try:
            size = _build_pdf(title, content, path)
        except Exception as exc:
            return ToolResult(f"failed to generate pdf at {raw_path}: {exc}", is_error=True)

        return ToolResult(f"generated pdf ({size} bytes) at {raw_path}")
