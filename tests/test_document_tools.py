"""Tests for document generation tools and markdown parser."""

from __future__ import annotations

import re
import zlib
from pathlib import Path

import docx
import pytest

from workbench.config import ToolsConfig
from workbench.core.protocols import ToolContext
from workbench.tools.documents import GenerateDocxTool, GeneratePdfTool
from workbench.tools.markdown_parser import (
    CodeBlock,
    Heading,
    HorizontalRule,
    InlineSpan,
    ListBlock,
    Paragraph,
    Table,
    parse_inlines,
    parse_markdown,
)


def _context(workspace_root: Path) -> ToolContext:
    async def _confirm(message: str) -> bool:
        return True

    return ToolContext(workspace_root=workspace_root, confirm=_confirm)


def test_markdown_parser_headings() -> None:
    text = "# Title\n## Subtitle\n### Section\n#### Sub-section\n##### Deep\n###### Deeper"
    blocks = parse_markdown(text)
    assert len(blocks) == 6
    assert blocks[0] == Heading(level=1, text="Title")
    assert blocks[1] == Heading(level=2, text="Subtitle")
    assert blocks[2] == Heading(level=3, text="Section")
    assert blocks[3] == Heading(level=4, text="Sub-section")
    assert blocks[4] == Heading(level=5, text="Deep")
    assert blocks[5] == Heading(level=6, text="Deeper")


def test_markdown_parser_code_blocks() -> None:
    text = "```python\ndef hello():\n    return 'world'\n```"
    blocks = parse_markdown(text)
    assert len(blocks) == 1
    assert isinstance(blocks[0], CodeBlock)
    assert blocks[0].language == "python"
    assert blocks[0].code == "def hello():\n    return 'world'"


def test_markdown_parser_lists() -> None:
    text = "- item 1\n- item 2\n- item 3\n\n1. first\n2. second\n3. third"
    blocks = parse_markdown(text)
    assert len(blocks) == 2
    assert blocks[0] == ListBlock(ordered=False, items=["item 1", "item 2", "item 3"])
    assert blocks[1] == ListBlock(ordered=True, items=["first", "second", "third"])


def test_markdown_parser_tables() -> None:
    text = "| Col A | Col B |\n| --- | --- |\n| 1 | 2 |\n| 3 | 4 |"
    blocks = parse_markdown(text)
    assert len(blocks) == 1
    assert isinstance(blocks[0], Table)
    assert blocks[0].headers == ["Col A", "Col B"]
    assert blocks[0].rows == [["1", "2"], ["3", "4"]]


def test_markdown_parser_horizontal_rule() -> None:
    text = "Intro\n\n---\n\nOutro"
    blocks = parse_markdown(text)
    assert len(blocks) == 3
    assert isinstance(blocks[0], Paragraph)
    assert isinstance(blocks[1], HorizontalRule)
    assert isinstance(blocks[2], Paragraph)


def test_markdown_parser_inlines() -> None:
    spans = parse_inlines("Hello **bold** and *italic* and `code` world")
    assert all(isinstance(s, InlineSpan) for s in spans)
    assert any(s.text == "bold" and s.bold for s in spans)
    assert any(s.text == "italic" and s.italic for s in spans)
    assert any(s.text == "code" and s.code for s in spans)


def test_markdown_parser_mixed() -> None:
    text = """# Executive Summary

This is the introductory paragraph with **bold** text.

## Findings

- High performance
- Low memory usage

| Metric | Score |
|---|---|
| Latency | 5ms |
| Memory | 250MB |

```json
{"status": "ok"}
```

---

Final remarks.
"""
    blocks = parse_markdown(text)
    types = [type(b) for b in blocks]
    assert types == [
        Heading,
        Paragraph,
        Heading,
        ListBlock,
        Table,
        CodeBlock,
        HorizontalRule,
        Paragraph,
    ]


@pytest.mark.asyncio
async def test_generate_docx_basic(tmp_path: Path) -> None:
    tool = GenerateDocxTool(ToolsConfig())
    ctx = _context(tmp_path)
    res = await tool.execute(
        {
            "title": "Quarterly Report",
            "content": "# Overview\nThis is a report.\n\n- Point 1\n- Point 2",
            "path": "report.docx",
        },
        ctx,
    )
    assert not res.is_error
    assert "generated docx" in res.content
    target = tmp_path / "report.docx"
    assert target.exists()

    doc = docx.Document(str(target))
    headings = [
        p.text
        for p in doc.paragraphs
        if p.text.startswith("Quarterly Report") or p.text.startswith("Overview")
    ]
    assert "Quarterly Report" in headings
    assert "Overview" in headings


@pytest.mark.asyncio
async def test_generate_docx_with_code_and_table(tmp_path: Path) -> None:
    tool = GenerateDocxTool(ToolsConfig())
    ctx = _context(tmp_path)
    content = "```python\nx = 42\n```\n\n| H1 | H2 |\n|---|---|\n| A | B |"
    res = await tool.execute(
        {"title": "Tech Specs", "content": content, "path": "specs.docx"}, ctx
    )
    assert not res.is_error
    target = tmp_path / "specs.docx"
    assert target.exists()

    doc = docx.Document(str(target))
    assert len(doc.tables) == 1
    assert doc.tables[0].rows[0].cells[0].text == "H1"
    assert doc.tables[0].rows[1].cells[0].text == "A"


@pytest.mark.asyncio
async def test_generate_docx_path_jail(tmp_path: Path) -> None:
    tool = GenerateDocxTool(ToolsConfig())
    ctx = _context(tmp_path)
    res = await tool.execute(
        {"title": "Escape", "content": "test", "path": "../escape.docx"}, ctx
    )
    assert res.is_error
    assert res.blocked_reason == "path_jail"


@pytest.mark.asyncio
async def test_generate_docx_bad_extension(tmp_path: Path) -> None:
    tool = GenerateDocxTool(ToolsConfig())
    ctx = _context(tmp_path)
    res = await tool.execute(
        {"title": "Doc", "content": "test", "path": "document.txt"}, ctx
    )
    assert res.is_error
    assert "must end with .docx" in res.content


@pytest.mark.asyncio
async def test_generate_pdf_basic(tmp_path: Path) -> None:
    tool = GeneratePdfTool(ToolsConfig())
    ctx = _context(tmp_path)
    res = await tool.execute(
        {
            "title": "Architecture Overview",
            "content": (
                "# Introduction\nOur system architecture is modular.\n\n"
                "- Component A\n- Component B"
            ),
            "path": "architecture.pdf",
        },
        ctx,
    )
    assert not res.is_error
    assert "generated pdf" in res.content
    target = tmp_path / "architecture.pdf"
    assert target.exists()

    raw = target.read_bytes()
    assert raw.startswith(b"%PDF")
    decompressed = b""
    for s in re.findall(b"stream\r?\n(.*?)\r?\nendstream", raw, re.DOTALL):
        try:
            decompressed += zlib.decompress(s)
        except Exception:
            decompressed += s
    assert b"Architecture Overview" in decompressed
    assert b"Introduction" in decompressed
    assert b"modular" in decompressed


@pytest.mark.asyncio
async def test_generate_pdf_with_code_and_table(tmp_path: Path) -> None:
    tool = GeneratePdfTool(ToolsConfig())
    ctx = _context(tmp_path)
    content = "```python\nprint('hello world')\n```\n\n| Key | Value |\n|---|---|\n| Alpha | 100 |"
    res = await tool.execute(
        {"title": "Report", "content": content, "path": "out/report.pdf"}, ctx
    )
    assert not res.is_error
    target = tmp_path / "out" / "report.pdf"
    assert target.exists()

    raw = target.read_bytes()
    assert raw.startswith(b"%PDF")
    decompressed = b""
    for s in re.findall(b"stream\r?\n(.*?)\r?\nendstream", raw, re.DOTALL):
        try:
            decompressed += zlib.decompress(s)
        except Exception:
            decompressed += s
    assert b"hello world" in decompressed
    assert b"Alpha" in decompressed


@pytest.mark.asyncio
async def test_generate_pdf_path_jail(tmp_path: Path) -> None:
    tool = GeneratePdfTool(ToolsConfig())
    ctx = _context(tmp_path)
    res = await tool.execute(
        {"title": "Escape", "content": "test", "path": "../escape.pdf"}, ctx
    )
    assert res.is_error
    assert res.blocked_reason == "path_jail"


@pytest.mark.asyncio
async def test_generate_pdf_bad_extension(tmp_path: Path) -> None:
    tool = GeneratePdfTool(ToolsConfig())
    ctx = _context(tmp_path)
    res = await tool.execute(
        {"title": "Doc", "content": "test", "path": "document.docx"}, ctx
    )
    assert res.is_error
    assert "must end with .pdf" in res.content
